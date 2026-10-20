from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# 创建演示文稿 16:9
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 颜色
RUST = RGBColor(0xCE, 0x42, 0x2B)
DARK = RGBColor(0x11, 0x11, 0x14)
DARK_CARD = RGBColor(0x1E, 0x1E, 0x24)
GRAY = RGBColor(0x9C, 0xA3, 0xAF)
LIGHT_GRAY = RGBColor(0xD1, 0xD5, 0xDB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x3B, 0x82, 0xF6)

def set_bg(slide, color=DARK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def text(slide, left, top, width, height, content, size=18, bold=False,
         color=WHITE, align=PP_ALIGN.LEFT, font="Microsoft YaHei", wrap=True):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.text = content
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    p.alignment = align
    return box

def line(slide, left, top, width, color=RUST, thickness=3):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                    Inches(width), Inches(thickness/72))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def card(slide, left, top, width, height, bg=DARK_CARD, radius=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                    Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.fill.background()
    try:
        shape.adjustments[0] = radius
    except:
        pass
    return shape

def badge(slide, left, top, content, bg=RUST):
    w = min(max(len(content) * 0.15 + 0.3, 1.1), 3.0)
    card(slide, left, top, w, 0.38, bg)
    text(slide, left, top + 0.04, w, 0.3, content, size=12, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER)
    return w

def big_num(slide, left, top, num, label, sub=""):
    text(slide, left, top, 2.8, 0.8, num, size=42, bold=True, color=RUST)
    text(slide, left, top + 0.75, 2.8, 0.35, label, size=14, color=LIGHT_GRAY)
    if sub:
        text(slide, left, top + 1.05, 2.8, 0.3, sub, size=11, color=GRAY)

def bullets(slide, left, top, width, height, items, size=15, color=LIGHT_GRAY, spacing=0.7):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Microsoft YaHei"
        p.space_after = Pt(spacing * 12)
    return box

# ===== 第1页：封面 =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

# 左侧装饰条
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                          Inches(0.12), Inches(7.5))
bar.fill.solid(); bar.fill.fore_color.rgb = RUST; bar.line.fill.background()

# 右上角橙色块
deco = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.8), Inches(-0.5),
                           Inches(3.0), Inches(2.8))
deco.fill.solid(); deco.fill.fore_color.rgb = RUST; deco.line.fill.background()
deco.rotation = -12

text(s, 0.8, 2.0, 10, 1.2, "周景潇", size=64, bold=True)
badge(s, 0.8, 3.35, "AI Agent 开发")
badge(s, 3.0, 3.35, "后端系统开发", DARK_CARD)
badge(s, 5.2, 3.35, "智能体平台研发", DARK_CARD)

text(s, 0.8, 4.15, 11, 0.7,
     "用 Rust 构建本地优先、零外部依赖的 AI Agent 基础设施",
     size=22, color=GRAY)

line(s, 0.8, 5.7, 8, GRAY, 1)
text(s, 0.8, 5.9, 11, 0.5,
     "甘肃农业大学 · 数据科学与大数据技术 · 本科 2027 届",
     size=15, color=GRAY)
text(s, 0.8, 6.4, 11, 0.5,
     "13626566112  ·  2241470466@qq.com  ·  github.com/juice094",
     size=14, color=LIGHT_GRAY)

# ===== 第2页：我是谁 =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0.8, 0.5, 10, 0.8, "我是谁", size=36, bold=True)
line(s, 0.8, 1.25, 1.5, RUST, 3)

text(s, 0.8, 1.8, 7.5, 2.4,
     "我不是在写‘又一个本地大模型客户端’，而是在构建让 AI 身份能够跨设备、跨模型、跨网络边界保持连续的分布式基础设施。\n\n"
     "三个独立 Rust 系统、40+ workspace crate、全绿测试、零生产 panic，"
     "工程上坚持：本地优先 · 零外部依赖 · 纯 Rust 全栈。",
     size=18, color=LIGHT_GRAY)

big_num(s, 9.3, 1.8, "3", "系统级项目", "production-grade Rust systems")
big_num(s, 9.3, 3.5, "40+", "workspace crates")
big_num(s, 9.3, 5.2, "0", "生产 panic", "unwrap / expect / panic free")

# ===== 第3页：技术主线 =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0.8, 0.5, 10, 0.8, "技术主线", size=36, bold=True)
line(s, 0.8, 1.25, 1.5, RUST, 3)

# 三个主线卡片
main_lines = [
    ("本地优先", "用户数据不强制上云，推理和记忆发生在设备本地，隐私可控。",
     "Clarity SQLite 记忆 · devbase 本地 embedding"),
    ("零外部依赖", "部署步骤最小化，单二进制分发，避免 Python/Node/向量数据库依赖。",
     "devbase 4→1 部署步骤 · Clarity cargo install 即运行"),
    ("纯 Rust 全栈", "从协议编解码到 UI 到移动端 FFI 全用 Rust，统一工具链和性能基线。",
     "BEP 协议 · MCP 协议 · Agent loop"),
]

for i, (title, desc, example) in enumerate(main_lines):
    x = 0.7 + i * 4.2
    card(s, x, 1.7, 3.9, 4.8)
    text(s, x + 0.2, 1.9, 3.5, 0.6, title, size=22, bold=True)
    text(s, x + 0.2, 2.6, 3.5, 1.6, desc, size=15, color=LIGHT_GRAY)
    text(s, x + 0.2, 4.4, 3.5, 1.6, f"例如：{example}", size=13, color=GRAY)

# ===== 第4页：Clarity =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0.8, 0.5, 9, 0.8, "Clarity", size=40, bold=True)
text(s, 0.8, 1.15, 10, 0.5,
     "面向具身智能的本地优先分布式人格连续运行时", size=18, color=GRAY)
badge(s, 10.5, 0.6, "核心项目")

# 问题与定位
text(s, 0.8, 1.9, 11.7, 0.8,
     "现有 AI 助手在不同终端启动时会面临上下文隔离、角色断裂、跨日连续性问题。"
     "Clarity 让同一个 AI 身份在不同模型、设备、网络边界之间保持人格、记忆和社会关系一致。",
     size=17, color=LIGHT_GRAY)

# 四层模型
layers = [
    ("Identity", "身份认同：我是谁、价值观、社会关系"),
    ("Personality", "人格配置：语气、风格、伦理边界"),
    ("Memory Graph", "记忆图谱：事实、情感、关系、项目上下文"),
    ("Avatar", "分身实例：运行在具体设备上的模型实例"),
]
for i, (title, desc) in enumerate(layers):
    y = 3.0 + i * 0.92
    card(s, 0.8, y, 0.4, 0.4, RUST)
    text(s, 0.9, y + 0.02, 0.2, 0.35, str(i+1), size=14, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER)
    text(s, 1.4, y, 4.5, 0.4, title, size=16, bold=True)
    text(s, 1.4, y + 0.38, 4.5, 0.4, desc, size=13, color=GRAY)

# 右侧数据
big_num(s, 7.5, 3.0, "23", "workspace crates")
big_num(s, 7.5, 4.4, "152K", "行 Rust 代码", "LOC")
big_num(s, 7.5, 5.8, "1,243", "测试", "零 Clippy warning")

# 关键决策
bullets(s, 10.2, 3.0, 2.8, 3.8,
        ["ReAct + Plan 双模式", "MCP 三传输", "SQLite 向量 UDF", "ChaCha20-Poly1305 加密"],
        size=13, color=GRAY)

# ===== 第5页：syncthing-rust =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0.8, 0.5, 10, 0.8, "syncthing-rust", size=40, bold=True)
text(s, 0.8, 1.15, 10, 0.5,
     "与 Go 版 Syncthing wire-compatible 的 BEP 协议 Rust 实现", size=18, color=GRAY)

# 核心能力
text(s, 0.8, 1.9, 7, 1.0,
     "从零实现 Block Exchange Protocol，用 prost 做 Protobuf 编解码，"
     "rustls 实现 TLS 1.3 传输，和官方 Go 版真实互通。",
     size=17, color=LIGHT_GRAY)

# 架构层次
arch = [
    ("协议层", "BEP wire format"),
    ("传输层", "TLS 1.3 + TCP/UDP + NAT 穿透"),
    ("同步层", "块级 delta 同步"),
    ("存储层", "Block index + 版本归档"),
]
for i, (title, desc) in enumerate(arch):
    x = 0.8 + i * 3.1
    card(s, x, 3.0, 2.8, 1.4)
    text(s, x + 0.15, 3.15, 2.5, 0.4, title, size=16, bold=True)
    text(s, x + 0.15, 3.6, 2.5, 0.7, desc, size=13, color=GRAY)

# 数据
metrics = [("8", "crates"), ("5", "binaries"), ("59K", "LOC"), ("~12MB", "static binary")]
for i, (num, label) in enumerate(metrics):
    x = 0.8 + i * 3.1
    text(s, x, 4.8, 2.8, 0.7, num, size=30, bold=True, color=RUST, align=PP_ALIGN.CENTER)
    text(s, x, 5.4, 2.8, 0.4, label, size=13, color=GRAY, align=PP_ALIGN.CENTER)

# 亮点
bullets(s, 0.8, 6.0, 12, 1.0,
        ["多路径 NAT 穿透：UDP 广播 / STUN / UPnP / Relay", "块级 delta 同步，只传变更块",
         "Windows 句柄冲突退避 + 三路文本合并 + 版本归档", "Tailscale 验证跨国组网稳定性"],
        size=14, color=LIGHT_GRAY, spacing=0.5)

# ===== 第6页：devbase =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0.8, 0.5, 10, 0.8, "devbase", size=40, bold=True)
text(s, 0.8, 1.15, 10, 0.5,
     "Developer Workspace Compiler — 把代码和笔记编译成 AI 能理解的上下文", size=18, color=GRAY)

# Compiler 隐喻
text(s, 0.8, 1.9, 11.7, 0.8,
     "像编译器把源码转成机器码一样，devbase 把分散的 Git 仓库、Markdown 笔记、Skill 脚本和 YAML 工作流，"
     "编译成结构化的代码上下文、知识记忆和 Agent 推理路径。",
     size=17, color=LIGHT_GRAY)

# 三层架构
layers2 = [
    ("感知层", "tree-sitter 多语言 AST\n代码符号解析"),
    ("知识层", "SQLite 图存储\nBM25 + embedding 检索"),
    ("策略层", "YAML DAG 工作流\nSkill 编排"),
]
for i, (title, desc) in enumerate(layers2):
    x = 0.8 + i * 4.2
    card(s, x, 3.0, 3.9, 2.2)
    text(s, x + 0.2, 3.15, 3.5, 0.5, title, size=18, bold=True)
    text(s, x + 0.2, 3.75, 3.5, 1.2, desc, size=14, color=GRAY)
    if i < 2:
        line(s, x + 3.95, 3.9, 0.3, GRAY, 2)

# 数据
big_num(s, 0.8, 5.6, "12", "workspace crates")
big_num(s, 4.0, 5.6, "71", "MCP 工具")
big_num(s, 7.2, 5.6, "56K", "LOC")
big_num(s, 10.4, 5.6, "8.7MB", "二进制", "single static binary")

# ===== 第7页：项目关联 =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0.8, 0.5, 10, 0.8, "三个项目如何互相支撑", size=36, bold=True)
line(s, 0.8, 1.25, 3, RUST, 3)

# 项目节点
nodes = [
    ("Clarity", "AI Agent 运行时\n人格 · 记忆 · 工具", 0.8, 2.5, RUST),
    ("syncthing-rust", "P2P 同步协议\n文件 · 数据 · 发现", 4.8, 2.5, BLUE),
    ("devbase", "工作空间编译器\n代码 · 笔记 · 工作流", 8.8, 2.5, RGBColor(0x10, 0xB9, 0x81)),
]
for title, desc, x, y, color in nodes:
    card(s, x, y, 3.6, 1.8, color)
    text(s, x + 0.15, y + 0.15, 3.3, 0.45, title, size=18, bold=True)
    text(s, x + 0.15, y + 0.65, 3.3, 1.0, desc, size=13, color=WHITE)

# 连接线（用细矩形模拟）
line(s, 4.45, 3.35, 0.4, GRAY, 2)
line(s, 8.45, 3.35, 0.4, GRAY, 2)

# 关联说明
relations = [
    "devbase 复用 syncthing-rust → 实现 P2P workspace 同步",
    "Clarity 与 devbase 都基于 MCP 协议 → 统一工具调用标准",
    "三个项目共享同一工程基线 → SQLite 内嵌 · 零外部依赖 · 纯 Rust",
]
bullets(s, 1.5, 4.8, 10.3, 1.8, relations, size=16, color=LIGHT_GRAY)

# ===== 第8页：研究经历 =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0.8, 0.5, 10, 0.8, "研究经历", size=36, bold=True)
line(s, 0.8, 1.25, 1.5, RUST, 3)

text(s, 0.8, 1.7, 11, 0.6,
     "检索增强生成中输出结构的实证研究", size=24, bold=True)
text(s, 0.8, 2.3, 8, 0.4, "github.com/juice094/acr-select", size=14, color=BLUE)

text(s, 0.8, 2.9, 8, 1.6,
     "RAG 系统的输出质量不仅取决于检索准确性，还受输出格式（JSON / Markdown / 自然语言）"
     "与检索深度的交互影响。本研究设计多因子受控实验，量化评估这种交互效应。",
     size=17, color=LIGHT_GRAY)

# 数据卡片
data = [("40", "实验条件"), ("3000+", "样本"), ("4", "模型架构"), ("论文", "撰写中")]
for i, (num, label) in enumerate(data):
    x = 0.8 + i * 3.1
    card(s, x, 5.0, 2.8, 1.8)
    text(s, x, 5.25, 2.8, 0.8, num, size=36, bold=True, color=RUST, align=PP_ALIGN.CENTER)
    text(s, x, 6.05, 2.8, 0.4, label, size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ===== 第9页：其他 =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0.8, 0.5, 10, 0.8, "认证与其他", size=36, bold=True)
line(s, 0.8, 1.25, 2, RUST, 3)

items = [
    "认证：AI 应用工程师（中级）— 工业和信息化部认证",
    "语言：日语专业四级（CJT-4）· 大学英语四级（CET-4）",
    "到岗时间：2026.07；可实习 3–6 个月，每周 5 天",
    "公网可访问：Clarity Web 前端（Vercel）、syncthing-rust P2P 节点（日本 VPS）",
]
bullets(s, 0.8, 1.8, 11.7, 3.5, items, size=18, color=LIGHT_GRAY, spacing=1.0)

# ===== 第10页：结尾 =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)

text(s, 0, 2.0, 13.333, 1.0, "期待与你一起", size=36, bold=True,
     color=WHITE, align=PP_ALIGN.CENTER)
text(s, 0, 2.9, 13.333, 0.8, "构建下一代 AI Agent 基础设施",
     size=28, bold=True, color=RUST, align=PP_ALIGN.CENTER)

line(s, 5.0, 4.0, 3.333, GRAY, 1)

text(s, 0, 4.5, 13.333, 0.5,
     "到岗时间：2026.07  ·  可实习 3–6 个月，每周 5 天",
     size=16, color=GRAY, align=PP_ALIGN.CENTER)
text(s, 0, 5.2, 13.333, 0.5,
     "13626566112  ·  2241470466@qq.com  ·  github.com/juice094",
     size=15, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# 保存
out = r'C:\Users\22414\dev\juice094\resume\周景潇-求职综合展示-v3.pptx'
prs.save(out)
print(f"PPT 已生成：{out}")
