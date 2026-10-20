from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

# 创建演示文稿 - 16:9
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 颜色系统 - Rust 科技风
RUST = RGBColor(0xCE, 0x42, 0x2B)
RUST_LIGHT = RGBColor(0xE8, 0x6D, 0x5A)
DARK = RGBColor(0x11, 0x11, 0x14)
DARK_CARD = RGBColor(0x1E, 0x1E, 0x24)
GRAY = RGBColor(0x9C, 0xA3, 0xAF)
LIGHT_GRAY = RGBColor(0xD1, 0xD5, 0xDB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_BLUE = RGBColor(0x3B, 0x82, 0xF6)

def set_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text(slide, left, top, width, height, text, font_size=18, bold=False,
             color=WHITE, align=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return txBox

def add_rounded_rect(slide, left, top, width, height, fill_color, line_color=None, radius=0.1):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                    Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    # 调整圆角
    if shape.has_text_frame:
        shape.adjustments[0] = radius
    return shape

def add_line(slide, left, top, width, color, thickness=2):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(thickness / 72))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    return line

def add_card_with_title(slide, left, top, width, height, title, content, title_size=20, content_size=14):
    """添加带标题的卡片"""
    card = add_rounded_rect(slide, left, top, width, height, DARK_CARD)

    # 标题
    add_text(slide, left + 0.2, top + 0.2, width - 0.4, 0.5, title,
             font_size=title_size, bold=True, color=WHITE)

    # 内容
    txBox = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.7),
                                      Inches(width - 0.4), Inches(height - 0.9))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = content
    p.font.size = Pt(content_size)
    p.font.color.rgb = GRAY
    p.font.name = "Microsoft YaHei"
    p.line_spacing = 1.3
    return card

def add_badge(slide, left, top, text, color=RUST):
    """添加小标签"""
    # 估算宽度
    width = max(len(text) * 0.16 + 0.3, 1.0)
    shape = add_rounded_rect(slide, left, top, width, 0.38, color)
    add_text(slide, left, top + 0.05, width, 0.28, text,
             font_size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    return shape

def add_big_number(slide, left, top, number, label, unit=""):
    """添加大数字展示"""
    add_text(slide, left, top, 2.5, 0.9, number,
             font_size=44, bold=True, color=RUST, align=PP_ALIGN.LEFT)
    add_text(slide, left, top + 0.75, 2.5, 0.4, unit,
             font_size=14, color=GRAY, align=PP_ALIGN.LEFT)
    add_text(slide, left, top + 1.05, 2.5, 0.4, label,
             font_size=12, color=LIGHT_GRAY, align=PP_ALIGN.LEFT)

# ===== 第1页：封面 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, DARK)

# 左侧装饰条
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
                              Inches(0.12), Inches(7.5))
bar.fill.solid()
bar.fill.fore_color.rgb = RUST
bar.line.fill.background()

# 右上角装饰块
deco = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.5), Inches(0),
                               Inches(2.833), Inches(2.5))
deco.fill.solid()
deco.fill.fore_color.rgb = RUST
deco.line.fill.background()
deco.rotation = -15

# 姓名
add_text(slide, 0.8, 2.0, 10, 1.2, "周景潇", font_size=64, bold=True, color=WHITE)

# 岗位标签
add_badge(slide, 0.8, 3.35, "AI Agent 开发", RUST)
add_badge(slide, 3.0, 3.35, "后端系统开发", DARK_CARD)
add_badge(slide, 5.2, 3.35, "智能体平台研发", DARK_CARD)

# 标语
add_text(slide, 0.8, 4.1, 10, 0.8,
         "用 Rust 构建本地优先、零外部依赖的 AI Agent 系统",
         font_size=22, color=GRAY)

# 底部信息
add_line(slide, 0.8, 5.8, 8, GRAY, 1)
add_text(slide, 0.8, 6.0, 10, 0.5,
         "甘肃农业大学 · 数据科学与大数据技术 · 本科 2027 届",
         font_size=15, color=GRAY)
add_text(slide, 0.8, 6.5, 10, 0.5,
         "13626566112  ·  2241470466@qq.com  ·  github.com/juice094",
         font_size=14, color=LIGHT_GRAY)

# ===== 第2页：个人定位 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, DARK)

add_text(slide, 0.8, 0.6, 10, 0.8, "个人定位", font_size=36, bold=True, color=WHITE)
add_line(slide, 0.8, 1.35, 1.5, RUST, 3)

add_text(slide, 0.8, 1.8, 7, 2.5,
         "独立用 Rust 从零构建过三个系统级项目，覆盖 AI Agent 运行时、P2P 文件同步、开发者知识库。\n\n"
         "熟悉异步服务、嵌入式存储、网络协议和跨平台 UI，工程上追求零警告、零 panic、可单二进制分发。",
         font_size=18, color=LIGHT_GRAY)

# 右侧大数字
add_big_number(slide, 9.2, 1.8, "3", "独立系统级项目", "production-grade Rust systems")
add_big_number(slide, 9.2, 3.5, "44+", "workspace crates", "across all projects")
add_big_number(slide, 9.2, 5.2, "0", "生产 panic", "unwrap / expect / panic free")

# ===== 第3页：技术栈 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, DARK)

add_text(slide, 0.8, 0.5, 10, 0.8, "技术栈", font_size=36, bold=True, color=WHITE)
add_line(slide, 0.8, 1.25, 1.5, RUST, 3)

# 6个技术卡片
tech_cards = [
    ("异步与网络", "Tokio · Axum · Tower\nreqwest · WebSocket · rustls"),
    ("AI Agent", "ReAct / Plan 双模式\nMCP 协议 · Multi-Agent\nRAG 检索增强"),
    ("存储与检索", "SQLite · rusqlite\nTantivy · BM25\nembedding 向量搜索"),
    ("分布式与 P2P", "BEP 协议 · STUN\nUPnP · Relay\nNAT 穿透"),
    ("跨平台 UI", "egui · ratatui\nUniFFI · 移动端 FFI"),
    ("安全与工程化", "ChaCha20-Poly1305\nTLS 1.3 · Cargo Workspace\nGitHub Actions · Clippy"),
]

for i, (title, content) in enumerate(tech_cards):
    row = i // 3
    col = i % 3
    add_card_with_title(slide, 0.7 + col * 4.2, 1.6 + row * 2.7, 3.9, 2.3,
                        title, content, title_size=18, content_size=13)

# ===== 第4页：Clarity =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, DARK)

add_text(slide, 0.8, 0.5, 10, 0.8, "Clarity", font_size=40, bold=True, color=WHITE)
add_text(slide, 0.8, 1.15, 10, 0.5,
         "让 AI 助手跨设备保持一致人格的本地 Agent 系统",
         font_size=18, color=GRAY)
add_badge(slide, 10.5, 0.6, "核心项目", RUST)

# 左侧数据
add_big_number(slide, 0.8, 2.0, "21", "workspace crates")
add_big_number(slide, 0.8, 3.7, "150K", "行 Rust 代码", "LOC")
add_big_number(slide, 0.8, 5.4, "0", "生产 panic")

# 右侧描述卡片
desc = ("本地优先的 AI Agent 运行时，让同一个 AI 身份在手机、电脑、服务器上保持一致。\n\n"
        "• ReAct + Plan 双模式执行，步骤遗漏率从约 40% 降到接近零\n"
        "• MCP 工具协议三传输实现（stdio / SSE / WebSocket）\n"
        "• SQLite 单文件混合记忆检索：手写向量 UDF + BM25\n"
        "• 桌面 GUI / 终端 UI / 后台服务 / 移动端 FFI 共享同一内核")
add_card_with_title(slide, 4.2, 2.0, 8.3, 4.8, "", desc, title_size=1, content_size=16)

add_text(slide, 4.4, 6.6, 8, 0.4, "github.com/juice094/clarity · 2026.04 — 至今",
         font_size=13, color=ACCENT_BLUE)

# ===== 第5页：syncthing-rust =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, DARK)

add_text(slide, 0.8, 0.5, 10, 0.8, "syncthing-rust", font_size=40, bold=True, color=WHITE)
add_text(slide, 0.8, 1.15, 10, 0.5,
         "无需云盘的 P2P 文件同步工具",
         font_size=18, color=GRAY)

# 顶部数据条
stats = [("9+5", "lib + bin crates"), ("413", "测试通过"), ("16MB", "单二进制")]
for i, (num, label) in enumerate(stats):
    add_text(slide, 0.8 + i * 4.0, 2.0, 3.8, 0.8, num,
             font_size=32, bold=True, color=RUST, align=PP_ALIGN.LEFT)
    add_text(slide, 0.8 + i * 4.0, 2.7, 3.8, 0.4, label,
             font_size=13, color=GRAY, align=PP_ALIGN.LEFT)

add_line(slide, 0.8, 3.2, 11.7, DARK_CARD, 2)

# 描述
desc = ("官方 Go 版 Syncthing 的 Rust 重写，让多台设备直接互联同步文件。\n\n"
        "• 从零实现 BEP 协议，wire_compat 验证与 Go 版线路兼容\n"
        "• 多路径 NAT 穿透：UDP 广播 / STUN / UPnP / Relay\n"
        "• 块级同步 + Windows 句柄冲突退避 + 三路文本合并 + 版本归档")
add_card_with_title(slide, 0.8, 3.5, 11.7, 3.0, "", desc, title_size=1, content_size=16)

add_text(slide, 0.8, 6.6, 10, 0.4, "github.com/juice094/syncthing-rust · 2025.11 — 至今",
         font_size=13, color=ACCENT_BLUE)

# ===== 第6页：devbase =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, DARK)

add_text(slide, 0.8, 0.5, 10, 0.8, "devbase", font_size=40, bold=True, color=WHITE)
add_text(slide, 0.8, 1.15, 10, 0.5,
         "把代码和笔记编译成 AI 能理解的上下文",
         font_size=18, color=GRAY)

# 三个亮点卡片
highlights = [
    ("81", "MCP stdio 工具", "覆盖代码检索、Git 分析、笔记搜索、YAML 工作流"),
    ("0", "云端向量数据库依赖", "Tantivy + SQLite BLOB + 手写 cosine_similarity UDF"),
    ("5", "工作流步骤类型", "Skill / 子工作流 / 并行 / 条件判断 / 循环"),
]

for i, (num, title, content) in enumerate(highlights):
    card = add_rounded_rect(slide, 0.7 + i * 4.2, 2.0, 3.9, 3.8, DARK_CARD)
    add_text(slide, 0.9 + i * 4.2, 2.2, 3.5, 0.9, num,
             font_size=48, bold=True, color=RUST, align=PP_ALIGN.LEFT)
    add_text(slide, 0.9 + i * 4.2, 3.1, 3.5, 0.5, title,
             font_size=16, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
    add_text(slide, 0.9 + i * 4.2, 3.7, 3.5, 1.8, content,
             font_size=13, color=GRAY, align=PP_ALIGN.LEFT)

add_text(slide, 0.8, 6.6, 10, 0.4, "github.com/juice094/devbase · 2026.01 — 至今",
         font_size=13, color=ACCENT_BLUE)

# ===== 第7页：研究经历 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, DARK)

add_text(slide, 0.8, 0.5, 10, 0.8, "研究经历", font_size=36, bold=True, color=WHITE)
add_line(slide, 0.8, 1.25, 1.5, RUST, 3)

add_text(slide, 0.8, 1.7, 8, 0.6,
         "检索增强生成中输出结构的实证研究",
         font_size=24, bold=True, color=WHITE)
add_text(slide, 0.8, 2.3, 8, 0.4,
         "github.com/juice094/acr-select",
         font_size=14, color=ACCENT_BLUE)

add_text(slide, 0.8, 3.0, 8, 2.0,
         "围绕 RAG 系统中“输出格式如何影响内容质量”设计受控实验，"
         "量化评估输出结构与检索配置的交互效应。",
         font_size=17, color=LIGHT_GRAY)

# 数据展示
metrics = [("40", "实验条件"), ("3000+", "样本"), ("4", "模型架构")]
for i, (num, label) in enumerate(metrics):
    card = add_rounded_rect(slide, 0.8 + i * 4.0, 5.0, 3.5, 1.8, DARK_CARD)
    add_text(slide, 0.8 + i * 4.0, 5.2, 3.5, 0.9, num,
             font_size=40, bold=True, color=RUST, align=PP_ALIGN.CENTER)
    add_text(slide, 0.8 + i * 4.0, 6.0, 3.5, 0.5, label,
             font_size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ===== 第8页：结尾 =====
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, DARK)

# 中央大文字
add_text(slide, 0, 2.0, 13.333, 1.0, "期待与你一起", font_size=36, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, 0, 2.9, 13.333, 0.8, "构建下一代 AI Agent 基础设施",
         font_size=28, bold=True, color=RUST, align=PP_ALIGN.CENTER)

add_line(slide, 5.0, 4.0, 3.333, GRAY, 1)

add_text(slide, 0, 4.5, 13.333, 0.5,
         "到岗时间：2026.07  ·  可实习 3–6 个月，每周 5 天",
         font_size=16, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, 0, 5.2, 13.333, 0.5,
         "13626566112  ·  2241470466@qq.com  ·  github.com/juice094",
         font_size=15, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# 保存
output_path = r'C:\Users\22414\dev\juice094\resume\周景潇-求职综合展示-v2.pptx'
prs.save(output_path)
print(f"PPT 已生成：{output_path}")
